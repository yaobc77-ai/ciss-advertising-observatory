"""Freeze the supplied CLAIMS input lineage for private human source review.

Reads the original nested ZIP and pinned cleaning notebook as data. No upstream
code, model, network, database, classification or publication operation runs.
"""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import html
import io
import json
import re
import stat
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import quote, urlsplit

from openpyxl import load_workbook

from observatory.claims_linkage import (
    _ordered_occurrences,
    load_input_bytes,
    validate_retained_input_bytes,
)
from observatory.claims_projection import (
    UPSTREAM_NOTEBOOK_SHA256,
    UpstreamAsciiProjectionLocator,
    _fix_encoding,
    projection_contract,
    validate_projected_claim_span,
)
from observatory.claims_publication import load_audit
from observatory.claims_spans import ClaimSpanLocator, normalize_claim_text
from observatory.evaluate import canonical_digest, strict_json_loads

FORMAT = "claims-original-input-review-v1"
MAX_FILE_BYTES = 64 * 1024 * 1024
WORKBOOK_MEMBER = "New/native_ad_dataset.xlsx"
TEXT_PREFIX = "New/articletexts/"


def digest(data):
    return hashlib.sha256(data).hexdigest()


def encoded(value):
    return (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def read_bytes(path):
    with Path(path).open("rb") as stream:
        data = stream.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("input_too_large")
    return data


def notebook_settings(data):
    """Read three literal assignments; never execute notebook cells or imports."""
    if digest(data) != UPSTREAM_NOTEBOOK_SHA256:
        raise ValueError("notebook_differs_from_pinned_projection")
    document = strict_json_loads(data)
    settings = {}
    for cell in document["cells"]:
        if cell.get("cell_type") != "code":
            continue
        for node in ast.walk(ast.parse("".join(cell["source"]))):
            if not isinstance(node, ast.Assign) or len(node.targets) != 1:
                continue
            target = node.targets[0]
            if not isinstance(target, ast.Name) or target.id not in {
                "MIN_PARAGRAPH_LENGTH", "MIN_WORD_COUNT", "JUNK_PATTERNS",
            }:
                continue
            if target.id in settings:
                raise ValueError("duplicate_cleaning_setting")
            settings[target.id] = ast.literal_eval(node.value)
    if (type(settings.get("MIN_PARAGRAPH_LENGTH")) is not int
            or type(settings.get("MIN_WORD_COUNT")) is not int
            or not isinstance(settings.get("JUNK_PATTERNS"), list)
            or any(not isinstance(p, str) for p in settings["JUNK_PATTERNS"])):
        raise ValueError("literal_cleaning_settings_required")
    return settings


def historical_cleaning(text, settings):
    """Reproduce the pinned historical selection only; preserve raw text separately."""
    projected = _fix_encoding(text.replace("\r\n", "\n").replace("\r", "\n"))
    paragraphs = [normalize_claim_text(p) for p in re.split(r"\n\s*\n+", projected)]
    paragraphs = [p for p in paragraphs if p]
    junk = re.compile("|".join(settings["JUNK_PATTERNS"]), re.IGNORECASE)
    kept = [p for p in paragraphs if not (
        len(p) < settings["MIN_PARAGRAPH_LENGTH"]
        or len(p.split()) < settings["MIN_WORD_COUNT"]
        or junk.search(p)
        or sum(c.isalpha() for c in p) < len(p) * 0.5
        or p.startswith("http") or p.startswith("www.")
        or re.match(r"^by\s+[\w\s\.\-,]{2,50}\s*$", p, re.IGNORECASE)
    )]
    return paragraphs, kept


def safe_url(value):
    if not isinstance(value, str) or any(c.isspace() for c in value):
        return None
    try:
        parsed = urlsplit(value)
        if (parsed.scheme not in {"http", "https"} or not parsed.hostname
                or parsed.username is not None or parsed.password is not None):
            return None
        _ = parsed.port
    except ValueError:
        return None
    return value


def table_text(value):
    text = html.escape(str(value).replace("|", " / ").replace("\r", " ").replace("\n", " "), quote=False)
    return re.sub(r"([\\`*_{}\[\]()#+.!-])", r"\\\1", text)


def code_block(text):
    # A source containing fences must not terminate the raw-text block.
    longest = max((len(m.group()) for m in re.finditer(r"`+", text)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}text\n{text}\n{fence}"


def file_link(label, path):
    target = quote(Path(path).resolve().as_posix(), safe="/:")
    return f"[{label}](<{target}>)"


def read_manual_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (len(entries) > 5000 or len({p.filename for p in entries}) != len(entries)
                or sum(p.file_size for p in entries) > MAX_FILE_BYTES):
            raise ValueError("duplicate_or_oversize_archive_members")
        if any(p.flag_bits & 1 or stat.S_ISLNK(p.external_attr >> 16) for p in entries):
            raise ValueError("encrypted_or_symlink_archive_member")
        workbook = archive.read(WORKBOOK_MEMBER)
        sources = {}
        selected_names = {WORKBOOK_MEMBER}
        for item in entries:
            if item.is_dir() or not item.filename.startswith(TEXT_PREFIX):
                continue
            name = item.filename[len(TEXT_PREFIX):]
            if not re.fullmatch(r"(?:0|[1-9][0-9]*)\.txt", name):
                raise ValueError("canonical_article_text_filename_required")
            article_id = name[:-4]
            raw = archive.read(item)
            sources[article_id] = {"member": item.filename, "bytes": raw,
                                   "text": raw.decode("utf-8"), "encoding": "utf-8"}
            selected_names.add(item.filename)
        ignored = [p.filename for p in entries if not p.is_dir() and p.filename not in selected_names]
    if not sources:
        raise ValueError("original_article_texts_required")
    return workbook, sources, ignored


def read_metadata(data):
    # Bound the XLSX's internal ZIP as well as the enclosing source ZIP.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(p.file_size for p in archive.infolist()) > MAX_FILE_BYTES:
            raise ValueError("workbook_too_large")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError("explicit_original_metadata_sheet_required")
        sheet = workbook["Sheet1"]
        if sheet.max_row > 20000 or sheet.max_column > 100:
            raise ValueError("workbook_dimensions_too_large")
        rows = iter(sheet.iter_rows())
        headers = [cell.value for cell in next(rows)]
        if (any(not isinstance(h, str) for h in headers) or len(set(headers)) != len(headers)
                or not {"id", "url", "publisher", "title", "sponsor", "notes"} <= set(headers)):
            raise ValueError("original_metadata_headers_required")
        result = {}
        for number, cells in enumerate(rows, 2):
            if all(cell.value is None for cell in cells):
                continue
            values = {header: cell.value for header, cell in zip(headers, cells, strict=True)}
            identifier = values["id"]
            if type(identifier) is not int or identifier < 0 or str(identifier) in result:
                raise ValueError("unique_integer_metadata_id_required")
            formula_fields = [header for header, cell in zip(headers, cells, strict=True) if cell.data_type == "f"]
            values = {key: value.isoformat() if isinstance(value, (datetime, date)) else value
                      for key, value in values.items()}
            result[str(identifier)] = {"sheet": "Sheet1", "source_row": number,
                                      "values": values, "formula_fields": formula_fields}
    finally:
        workbook.close()
    return headers, result


def load_current_snapshot(data):
    value = strict_json_loads(data)
    rows = value["records"]
    by_url, seen = defaultdict(list), set()
    for row in rows:
        payload = row["payload"]
        if (row["record_id"] in seen or canonical_digest(payload) != row["version_id"]
                or payload["record_id"] != row["record_id"] or payload["dataset"] != row["dataset"]
                or row["body"] != payload["body"]
                or digest(row["body"].encode("utf-8")) != row["body_sha256"]):
            raise ValueError("current_source_identity_mismatch")
        seen.add(row["record_id"])
        if row["dataset"] == "native" and row["active"] is True:
            by_url[payload["url"]].append({key: row[key] for key in ("record_id", "version_id", "body_sha256")}
                                        | {key: payload[key] for key in ("url", "title", "publisher", "sponsor", "countable", "retrievable")})
    return by_url


def source_locations(text, paragraphs):
    strict = ClaimSpanLocator(text)
    projected = UpstreamAsciiProjectionLocator(text)
    rows, all_spans = [], []
    for paragraph in paragraphs:
        spans = {(span.start, span.end): span for span in strict.locate(paragraph.text)}
        if paragraph.text.isascii() and normalize_claim_text(paragraph.text) == paragraph.text:
            for span in projected.locate(paragraph.text):
                validate_projected_claim_span(text, paragraph.text, span)
                spans.setdefault((span.start, span.end), span)
        ordered = sorted(spans.values(), key=lambda s: (s.start, s.end, s.match_method))
        all_spans.append(ordered)
        rows.append({"input_id": paragraph.source_id, "csv_row": paragraph.csv_row,
                     "input_text": paragraph.text, "input_text_sha256": digest(paragraph.text.encode("utf-8")),
                     "locations": [{"start": s.start, "end": s.end, "quote": s.quote,
                                    "match_method": s.match_method, "lossy": getattr(s, "lossy", False),
                                    "context_before": text[max(0, s.start - 200):s.start],
                                    "context_after": text[s.end:s.end + 200]} for s in ordered]})
    return rows, _ordered_occurrences(all_spans)


def csv_bytes(rows):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]), lineterminator="\r\n")
    writer.writeheader()
    for row in rows:
        writer.writerow({key: "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))
                         else value for key, value in row.items()})
    return stream.getvalue().encode("utf-8-sig")


def write_packet(manual_zip, input_csv, cleaned_csv, notebook, output, *, current_snapshot=None, source_audit=None):
    output = Path(output)
    if output.exists() or output.is_symlink():
        raise ValueError("new_private_output_directory_required")
    paths = {"source_archive.zip": Path(manual_zip), "input_paragraphs.csv": Path(input_csv),
             "cleaned_paragraphs.csv": Path(cleaned_csv), "cleaning_notebook.ipynb": Path(notebook)}
    if current_snapshot is not None:
        paths["current_source_snapshot.json"] = Path(current_snapshot)
    frozen = {name: read_bytes(path) for name, path in paths.items()}
    settings = notebook_settings(frozen["cleaning_notebook.ipynb"])
    paragraphs = load_input_bytes(frozen["input_paragraphs.csv"])
    retained_check = validate_retained_input_bytes(paragraphs, frozen["cleaned_paragraphs.csv"])
    groups = defaultdict(list)
    for paragraph in paragraphs:
        groups[paragraph.article_id].append(paragraph)
    cleaned = defaultdict(list)
    for row in csv.DictReader(io.StringIO(frozen["cleaned_paragraphs.csv"].decode("utf-8-sig"), newline="")):
        cleaned[row["id"]].append(row["text"])
    workbook_bytes, sources, ignored = read_manual_archive(frozen["source_archive.zip"])
    headers, metadata = read_metadata(workbook_bytes)
    if set(metadata) != set(sources) or not set(groups) <= set(sources) or not set(cleaned) <= set(sources):
        raise ValueError("article_filename_and_metadata_id_sets_differ")
    current_by_url = load_current_snapshot(frozen["current_source_snapshot.json"]) if current_snapshot else {}
    audit = load_audit(source_audit) if source_audit is not None else None
    baseline, prior_text_candidates, audit_files = {}, defaultdict(dict), {}
    if audit is not None:
        expected = audit.report["input_file_sha256"]
        if any(expected.get(paths[name].name) != digest(frozen[name]) for name in
               ("input_paragraphs.csv", "cleaned_paragraphs.csv", "cleaning_notebook.ipynb")):
            raise ValueError("input_files_differ_from_prior_audit")
        for name, fingerprint in audit.artifact_hashes.items():
            data = read_bytes(audit.directory / name)
            if digest(data) != fingerprint:
                raise ValueError("prior_audit_artifact_changed")
            audit_files[f"prior_audit/{name}"] = data
        manifest_bytes = read_bytes(audit.directory / "audit_output_manifest.json")
        if digest(manifest_bytes) != audit.manifest_sha256:
            raise ValueError("prior_audit_manifest_changed")
        audit_files["prior_audit/audit_output_manifest.json"] = manifest_bytes
        baseline = {r["upstream_article_id"]: r for r in strict_json_loads(audit_files["prior_audit/claims_article_linkage.json"])}
        if set(baseline) != set(groups):
            raise ValueError("prior_audit_article_groups_differ")
        current_records = {r["record_id"]: r for rows in current_by_url.values() for r in rows}
        for row in strict_json_loads(audit_files["prior_audit/claims_source_linkage.json"]):
            for candidate in row["candidates"]:
                if candidate["current"] is not True:
                    continue
                key = (candidate["record_id"], candidate["version_id"])
                collection = prior_text_candidates[row["upstream_article_id"]]
                current = current_records.get(candidate["record_id"])
                if key not in collection:
                    version_matches = (current is not None and current["version_id"] == candidate["version_id"]
                                       and current["body_sha256"] == candidate["body_hash"] and current["url"] == candidate["url"])
                    collection[key] = {k: candidate[k] for k in ("record_id", "version_id", "body_hash", "url", "source_issues")}
                    collection[key].update({"matched_input_ids": [], "publisher": current["publisher"] if version_matches else None,
                                           "current_snapshot_version_matches": None if current_snapshot is None else version_matches})
                    if current is not None and not version_matches:
                        collection[key]["current_snapshot_candidate"] = current
                if row["source_result_id"] not in collection[key]["matched_input_ids"]:
                    collection[key]["matched_input_ids"].append(row["source_result_id"])
    inventory = []
    for article_id, source in sorted(sources.items(), key=lambda item: int(item[0])):
        raw_paragraphs, kept = historical_cleaning(source["text"], settings)
        inventory.append({"upstream_article_id": article_id, "member": source["member"],
                          "raw_bytes_sha256": digest(source["bytes"]), "encoding": source["encoding"],
                          "decoded_text_sha256": digest(source["text"].encode("utf-8")),
                          "source_nonempty_paragraphs": len(raw_paragraphs), "historical_kept_paragraphs": len(kept),
                          "reproduces_cleaned_sequence": kept == cleaned.get(article_id, []),
                          "provided_metadata": metadata[article_id]})
    inventory_by_id = {r["upstream_article_id"]: r for r in inventory}
    lineage, reviews = [], []
    for article_id, selected in sorted(groups.items(), key=lambda item: int(item[0])):
        source = sources[article_id]
        locations, order = source_locations(source["text"], selected)
        row = inventory_by_id[article_id]
        values = metadata[article_id]["values"]
        url = values["url"]
        current = current_by_url.get(url, []) if isinstance(url, str) and "url" not in metadata[article_id]["formula_fields"] else []
        prior = baseline.get(article_id)
        lineage.append(row | {"supplied_paragraphs": len(selected), "paragraphs": locations,
                              "all_retained_paragraphs_in_order": order,
                              "located_paragraphs": sum(bool(p["locations"]) for p in locations),
                              "current_exact_url_candidates": current, "prior_source_audit": prior,
                              "prior_current_text_candidates": list(prior_text_candidates[article_id].values()),
                              "url_identity_review": "pending", "classification_semantic_review": "pending",
                              "original_page_completeness": "unknown", "input_is_complete_original_article": False,
                              "source_association_approved": False, "publication_status": "hold"})
        reviews.append({"upstream_article_id": article_id, "title_as_supplied": values["title"],
                        "publisher_as_supplied": values["publisher"], "url_as_supplied": url,
                        "metadata_source_row": metadata[article_id]["source_row"],
                        "supplied_paragraphs": len(selected), "located_paragraphs": lineage[-1]["located_paragraphs"],
                        "all_retained_in_order": order, "current_exact_url_matches": len(current),
                        "raw_text_sha256": row["decoded_text_sha256"], "source_review": "pending",
                        "reviewer": "", "reviewed_at": "", "identity_notes": "",
                        "original_page_completeness": "unknown", "publication_status": "hold"})
    summary = {"schema_version": FORMAT, "prepared_at": datetime.now(timezone.utc).isoformat(),
               "source_file_sha256": {name: digest(data) for name, data in frozen.items()},
               "source_snapshot_scope": "supplied_frozen_file_not_live_database" if current_snapshot else "not_supplied",
               "workbook_member": WORKBOOK_MEMBER, "workbook_sha256": digest(workbook_bytes),
               "source_articles": len(sources), "metadata_records": len(metadata), "metadata_headers": headers,
               "ignored_archive_members": ignored, "retained_input_check": retained_check,
               "source_nonempty_paragraphs": sum(r["source_nonempty_paragraphs"] for r in inventory),
               "historical_kept_paragraphs": sum(r["historical_kept_paragraphs"] for r in inventory),
               "cleaned_sequences_reproduced": sum(r["reproduces_cleaned_sequence"] for r in inventory),
               "source_groups_with_kept_paragraphs": sum(r["historical_kept_paragraphs"] > 0 for r in inventory),
               "selected_article_groups": len(lineage), "selected_input_paragraphs": len(paragraphs),
               "selected_ordered_groups": sum(r["all_retained_paragraphs_in_order"] for r in lineage),
               "selected_groups_with_exact_current_url": sum(bool(r["current_exact_url_candidates"]) for r in lineage),
               "selected_unique_url_groups": sum(len(r["current_exact_url_candidates"]) == 1 for r in lineage),
               "selected_repeated_input_paragraphs": sum(len(p["locations"]) > 1 for r in lineage for p in r["paragraphs"]),
               "selected_original_span_candidates": sum(len(p["locations"]) for r in lineage for p in r["paragraphs"]),
               "prior_text_linkage": ({"all_retained_groups": sum(bool(r["current_record_candidates"]) for r in baseline.values()),
                                        "partial_groups": sum(not r["current_record_candidates"] and r["located_paragraphs"] > 0 for r in baseline.values()),
                                        "no_located_groups": sum(r["located_paragraphs"] == 0 for r in baseline.values())} if audit else None),
               "match_methods": dict(Counter(s["match_method"] for r in lineage for p in r["paragraphs"] for s in p["locations"])),
               "cleaning_settings": settings, "projection_contract": projection_contract(),
               "prior_audit_manifest_sha256": audit.manifest_sha256 if audit else None,
               "prior_eligible_classification_candidates": len(audit.candidates) if audit else None,
               "model_calls": 0, "network_calls": 0, "database_writes": 0,
               "approved_source_associations": 0, "results_published": 0,
               "reviewer": None, "classification_quality_scores": None}
    # Prepare everything before creating the output; a manifest is written last.
    artifacts = frozen | audit_files | {"source_workbook.xlsx": workbook_bytes, "source_inventory.json": encoded(inventory),
                          "article_lineage.json": encoded(lineage), "pending_source_review.csv": csv_bytes(reviews),
                          "preparation_summary.json": encoded(summary)}
    intro = ["# CLAIMS 原始输入来源审查", "",
             f"已保存 {len(lineage)} 个文章组、{len(paragraphs)} 个保留输入段落的原始文本和表格对应。",
             "文件名与表格 ID 是上游输入链路；URL、媒体、赞助方仍是原表提供的待审信息。当前库 URL 相同也不证明正文版本相同。",
             "原文保留全部字符与被旧流程过滤的内容；已定位、语义支持、事实成立及发布授权分别判断。全部审核待定，不发布分类。", "",
             "| 旧文章组 | 原表标题 | 保留段落位置 | 当前库精确 URL 候选 | 详情 |",
             "| --- | --- | ---: | ---: | --- |"]
    for row in lineage:
        article_id = row["upstream_article_id"]
        values = row["provided_metadata"]["values"]
        raw_path = output / "original" / f"{article_id}.txt"
        article_path = output / "articles" / f"{article_id}.md"
        artifacts[f"original/{article_id}.txt"] = sources[article_id]["bytes"]
        lines = [f"# 旧文章组 {article_id}", "", "来源、语义和发布全部待审；本文件不提供分类结论。", "",
                 f"原始 ZIP 成员：`{row['member']}`；表格 `Sheet1` 第 {row['provided_metadata']['source_row']} 行。",
                 "", "| 原表字段（未核实） | 原值 |", "| --- | --- |"]
        lines += [f"| {table_text(key)} | {table_text(value)} |" for key, value in values.items()]
        public_url = safe_url(values["url"]) if "url" not in row["provided_metadata"]["formula_fields"] else None
        if public_url:
            lines += ["", f"[原表 URL（未联网核验）](<{quote(public_url, safe=':/?#[]@!$&()*+,;=%')}>)"]
        lines += ["", "## 当前库与旧审计", "", code_block(json.dumps({
            "current_exact_url_candidates": row["current_exact_url_candidates"],
            "prior_source_audit": row["prior_source_audit"],
            "prior_current_text_candidates": row["prior_current_text_candidates"]}, ensure_ascii=False, indent=2)),
                  "", "## 原始保存文本", "", file_link("打开原始 TXT", raw_path), "",
                  "完整原页面、图片和视频覆盖尚未确认；旧清洗保留序列不等于整篇原文。", "", code_block(sources[article_id]["text"]),
                  "", "## 保留输入段落与定位", "",
                  "位置为原始 TXT 的 Python Unicode 字符半开区间；投影仅用于比对，原文不改。重复出现全部保留，不选其中一处当作人工判断。"]
        for paragraph in row["paragraphs"]:
            lines += ["", f"### 输入 {paragraph['input_id']}（CSV 第 {paragraph['csv_row']} 行）", "", code_block(paragraph["input_text"])]
            if not paragraph["locations"]:
                lines += ["", "未定位；不能据此宣布该文章不含主张。"]
            for location in paragraph["locations"]:
                lines += ["", f"原文位置 [{location['start']}, {location['end']})；{location['match_method']}；lossy={location['lossy']}。", "",
                          code_block(location["context_before"] + location["quote"] + location["context_after"])]
        artifacts[f"articles/{article_id}.md"] = ("\n".join(lines) + "\n").encode("utf-8")
        intro.append(f"| {article_id} | {table_text(values['title'])} | {row['located_paragraphs']}/{row['supplied_paragraphs']} | {len(row['current_exact_url_candidates'])} | {file_link('审查详情', article_path)} |")
    intro += ["", file_link("待审 CSV", output / "pending_source_review.csv"), "",
              "下一步：先确认文件与原表来源、当前正文准入和版本，再按分类表检查每个结果。此 CSV 不是 claims-import 的批准文件。审核人暂时待定。"]
    artifacts["index.md"] = ("\n".join(intro) + "\n").encode("utf-8")
    if any(read_bytes(path) != frozen[name] for name, path in paths.items()):
        raise ValueError("input_changed_during_preparation")
    if audit is not None and load_audit(source_audit).manifest_sha256 != audit.manifest_sha256:
        raise ValueError("prior_audit_changed_during_preparation")
    output.mkdir(parents=True, exist_ok=False)
    for name, data in artifacts.items():
        path = output / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
    manifest = {"schema_version": FORMAT, "complete": True,
                "artifacts": {name: digest(data) for name, data in artifacts.items()},
                "source_associations_approved": False, "labels_published": False}
    (output / "packet_manifest.json").write_bytes(encoded(manifest))
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("manual-zip", "input-csv", "cleaned-csv", "notebook", "output"):
        parser.add_argument("--" + name, required=True, type=Path)
    parser.add_argument("--current-snapshot", type=Path)
    parser.add_argument("--source-audit", type=Path)
    args = parser.parse_args()
    result = write_packet(args.manual_zip, args.input_csv, args.cleaned_csv, args.notebook, args.output,
                          current_snapshot=args.current_snapshot, source_audit=args.source_audit)
    print(json.dumps({key: result[key] for key in ("schema_version", "source_articles", "selected_article_groups",
                                                "selected_input_paragraphs", "selected_ordered_groups",
                                                "selected_groups_with_exact_current_url", "results_published")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
